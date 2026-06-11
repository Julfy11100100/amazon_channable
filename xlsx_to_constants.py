import ast
import sys
from pathlib import Path

from openpyxl import load_workbook

CONSTANTS_FILE = Path("constants.py")
VAR_NAME = "BARCODES"


def load_barcodes_from_constants(path: Path, var_name: str = VAR_NAME) -> dict[str, str]:
    source = path.read_text(encoding="utf-8")
    tree = ast.parse(source, filename=str(path))

    for node in tree.body:
        if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name):
            if node.target.id == var_name:
                value = ast.literal_eval(node.value)
                if not isinstance(value, dict):
                    raise ValueError(f"{var_name} должен быть dict")
                return {str(k): str(v) for k, v in value.items()}

        elif isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id == var_name:
                    value = ast.literal_eval(node.value)
                    if not isinstance(value, dict):
                        raise ValueError(f"{var_name} должен быть dict")
                    return {str(k): str(v) for k, v in value.items()}

    raise KeyError(f"Переменная {var_name!r} не найдена в {path}")


def load_barcodes_from_xlsx(path: Path) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    result: dict[str, str] = {}

    for barcode_cell, key_cell in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if barcode_cell is None or key_cell is None:
            continue

        barcode = str(barcode_cell).strip()
        key = str(key_cell).strip()

        if not barcode or not key:
            continue

        result[key] = barcode

    wb.close()
    return result


def render_barcodes_dict(data: dict[str, str], var_name: str = VAR_NAME) -> str:
    lines = [f"{var_name}: dict[str, str] = {{"]

    for key, value in data.items():
        lines.append(f'    "{key}": "{value}",')

    lines.append("}")
    lines.append("")
    return "\n".join(lines)


def replace_barcodes_block(source: str, new_block: str, var_name: str = VAR_NAME) -> str:
    tree = ast.parse(source)

    for node in tree.body:
        if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name) and node.target.id == var_name:
            start = node.lineno - 1
            end = node.end_lineno
            lines = source.splitlines()
            return "\n".join(lines[:start] + [new_block.rstrip()] + lines[end:]) + "\n"

        if isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id == var_name:
                    start = node.lineno - 1
                    end = node.end_lineno
                    lines = source.splitlines()
                    return "\n".join(lines[:start] + [new_block.rstrip()] + lines[end:]) + "\n"

    raise KeyError(f"Переменная {var_name!r} не найдена")


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage: python xlsx_to_constants.py <file.xlsx>")
        raise SystemExit(1)

    xlsx_path = Path(sys.argv[1])

    if not xlsx_path.exists():
        print(f"Файл не найден: {xlsx_path}")
        raise SystemExit(1)

    if not CONSTANTS_FILE.exists():
        print(f"Файл не найден: {CONSTANTS_FILE}")
        raise SystemExit(1)

    current = load_barcodes_from_constants(CONSTANTS_FILE)
    new_data = load_barcodes_from_xlsx(xlsx_path)

    current.update(new_data)

    source = CONSTANTS_FILE.read_text(encoding="utf-8")
    new_block = render_barcodes_dict(current)
    updated_source = replace_barcodes_block(source, new_block)

    CONSTANTS_FILE.write_text(updated_source, encoding="utf-8")

    print(f"Добавлено/обновлено: {len(new_data)} записей")


if __name__ == "__main__":
    main()
